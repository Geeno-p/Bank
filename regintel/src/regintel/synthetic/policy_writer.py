"""Render PolicySpec -> DOCX (Docling input) + Markdown (human-readable) + index.json."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt

from regintel.synthetic.spec import POLICIES, ControlSpec, Narrative, PolicySpec, SectionSpec

BANK = "Bharat Unity Bank Ltd."
CLASSIFICATION = "Internal | SYNTHETIC document for hackathon demonstration. Not a real bank policy."


def _key(no: str) -> tuple[int, ...]:
    return tuple(int(x) for x in no.split("."))


def _clauses(section: SectionSpec) -> list[Narrative | ControlSpec]:
    return sorted([*section.narrative, *section.controls], key=lambda x: _key(x.clause_no))


def _control_parts(c: ControlSpec) -> tuple[str, str, str]:
    label = f"{c.clause_no} {c.name} [{c.id}]."
    meta = f"Control owner: {c.owner_role}."
    if c.frequency:
        meta += f" Frequency: {c.frequency.value.replace('_', ' ').title()}."
    return label, c.text, meta


def _meta_rows(p: PolicySpec) -> list[tuple[str, str]]:
    return [
        ("Policy ID", p.id), ("Version", p.version), ("Policy owner", p.owner_role),
        ("Approved by", p.approved_by), ("Approved on", p.approved_on.isoformat()),
        ("Last reviewed", p.last_reviewed.isoformat()),
        ("Review frequency", p.review_frequency.value.replace("_", " ").title()),
        ("Classification", CLASSIFICATION),
    ]


def write_docx(p: PolicySpec, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)

    doc.add_heading(p.title, level=0)
    doc.add_paragraph(BANK)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in _meta_rows(p):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, v

    for s in p.sections:
        doc.add_heading(f"{s.no}. {s.title}", level=1)
        for item in _clauses(s):
            para = doc.add_paragraph()
            if isinstance(item, ControlSpec):
                label, text, meta = _control_parts(item)
                para.add_run(label + " ").bold = True
                para.add_run(text + " ")
                para.add_run(meta).italic = True
            else:
                para.add_run(f"{item.clause_no} ").bold = True
                para.add_run(item.text)

    for i, a in enumerate(p.annexures):
        letter = chr(ord("A") + i)
        doc.add_heading(a.title, level=1)
        for j, text in enumerate(a.paragraphs, 1):
            para = doc.add_paragraph()
            para.add_run(f"{letter}.{j} ").bold = True
            para.add_run(text)

    fixed = datetime(p.approved_on.year, p.approved_on.month, p.approved_on.day)
    cp = doc.core_properties
    cp.title, cp.author, cp.created, cp.modified = p.title, p.owner_role, fixed, fixed
    doc.save(path)


def write_markdown(p: PolicySpec, path: Path) -> None:
    lines = [f"# {p.title}", "", f"**{BANK}**", "", "| Field | Value |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in _meta_rows(p)]
    for s in p.sections:
        lines += ["", f"## {s.no}. {s.title}", ""]
        for item in _clauses(s):
            if isinstance(item, ControlSpec):
                label, text, meta = _control_parts(item)
                lines += [f"**{label}** {text} _{meta}_", ""]
            else:
                lines += [f"**{item.clause_no}** {item.text}", ""]
    for i, a in enumerate(p.annexures):
        letter = chr(ord("A") + i)
        lines += ["", f"## {a.title}", ""]
        lines += [f"**{letter}.{j}** {t}\n" for j, t in enumerate(a.paragraphs, 1)]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_policies(pol_dir: Path) -> list[dict]:
    index = []
    for p in POLICIES:
        write_docx(p, pol_dir / f"{p.id}.docx")
        write_markdown(p, pol_dir / f"{p.id}.md")
        index.append({
            "policy_id": p.id, "title": p.title, "version": p.version, "owner_role": p.owner_role,
            "approved_by": p.approved_by, "approved_on": p.approved_on.isoformat(),
            "last_reviewed": p.last_reviewed.isoformat(),
            "review_frequency": p.review_frequency.value, "topics": p.topics,
            "docx": f"policies/{p.id}.docx", "markdown": f"policies/{p.id}.md",
        })
    (pol_dir / "index.json").write_text(json.dumps(index, indent=2), encoding="utf-8")
    return index