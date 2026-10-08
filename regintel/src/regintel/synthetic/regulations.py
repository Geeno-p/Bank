"""SYNTHETIC regulation excerpts (paraphrased, NOT official text) for the demo when official PDFs
are unavailable. Clause numbers are illustrative.   python -m regintel.synthetic.regulations"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from regintel.settings import get_settings

DISCLAIMER = ("SYNTHETIC DEMO EXCERPT: paraphrased for a hackathon demonstration. Not the official text; "
              "clause numbering is illustrative. Verify against the official publication.")

REGS: dict[str, tuple[str, list[str]]] = {
    "RBI-KYC-MD": ("Know Your Customer (KYC) Directions - Demo Excerpt", [
        "CHAPTER I Preliminary",
        "1. These Directions apply to every Regulated Entity (RE), including scheduled commercial banks, "
        "unless stated otherwise.",
        "2. Terms such as Officially Valid Document and Principal Officer carry their meaning under the PML Rules.",
        "CHAPTER II Governance",
        "3. Every RE shall have a KYC policy duly approved by its Board of Directors.",
        "4. Every RE shall nominate a Designated Director and a Principal Officer and communicate their "
        "names and contact details to FIU-IND.",
        "CHAPTER III Customer Acceptance and Identification",
        "5. No RE shall open an account in an anonymous or fictitious or benami name.",
        "6. Every RE shall undertake customer due diligence using an Officially Valid Document or Aadhaar "
        "e-KYC authentication.",
        "7. Where the customer is a legal entity, the RE shall identify the beneficial owner and take "
        "reasonable measures to verify the identity of the beneficial owner.",
        "8. Before opening an account, the RE shall verify that the customer does not appear in the UNSC "
        "sanctions lists or the lists of individuals and entities designated under the UAPA.",
        "9. Video-based Customer Identification Process (V-CIP) shall be conducted only by a trained official "
        "of the RE, with end-to-end recording, geo-tagging of the customer and a liveness check.",
        "10. Every RE shall upload the KYC records of individual customers to the Central KYC Records "
        "Registry (CKYCR) within ten days of commencement of an account-based relationship.",
        "CHAPTER IV Risk Management",
        "11. Every RE shall carry out risk categorisation of customers as low, medium or high risk.",
        "12. Periodic updation of KYC shall be carried out as under:",
        "(a) at least once in every two years for high risk customers;",
        "(b) at least once in every eight years for medium risk customers;",
        "(c) at least once in every ten years for low risk customers.",
        "CHAPTER V Reporting and Record Keeping",
        "13. Every RE shall report suspicious transactions to FIU-IND within seven working days of arriving "
        "at a conclusion that a transaction is suspicious.",
        "14. Every RE shall furnish cash transaction reports to FIU-IND for every month by the 15th day of "
        "the succeeding month.",
        "15. Every RE shall preserve records of transactions and of identification of customers for at least "
        "five years after the business relationship has ended or the account has been closed, whichever is later.",
        "CHAPTER VI Specific Entities",
        "16. Payments Banks shall not open accounts for customers whose end-of-day balance would exceed the "
        "limit prescribed for Payments Banks.",
        "17. The Reserve Bank may issue clarifications on these Directions from time to time.",
    ]),
    "DPDP-2023": ("Digital Personal Data Protection - Demo Excerpt", [
        "CHAPTER I Preliminary",
        "1. This excerpt applies to the processing of digital personal data within India.",
        "CHAPTER II Obligations of Data Fiduciary",
        "2. A Data Fiduciary shall process personal data only for a lawful purpose for which the Data "
        "Principal has given consent or for certain legitimate uses.",
        "3. Every request for consent shall be accompanied or preceded by a notice describing the personal "
        "data and the purpose of processing.",
        "4. The Data Fiduciary shall give the Data Principal the option to access the notice in English or "
        "any language specified in the Eighth Schedule to the Constitution.",
        "5. Consent shall be free, specific, informed, unconditional and unambiguous with a clear affirmative action.",
        "6. A Data Fiduciary shall enable the Data Principal to withdraw consent at any time, with the ease "
        "of doing so comparable to the ease with which consent was given.",
        "7. A Data Fiduciary shall protect personal data by taking reasonable security safeguards to prevent "
        "personal data breach.",
        "8. In the event of a personal data breach, the Data Fiduciary shall intimate the Data Protection "
        "Board and each affected Data Principal.",
        "9. A Data Fiduciary shall erase personal data when the purpose is no longer being served or consent "
        "is withdrawn, unless retention is necessary for compliance with any law.",
        "10. A Data Fiduciary shall publish the business contact information of a person who is able to "
        "answer questions of the Data Principal about the processing of personal data.",
        "11. Before processing personal data of a child, a Data Fiduciary shall obtain verifiable consent of "
        "the parent or lawful guardian.",
        "12. A Data Fiduciary shall not undertake tracking, behavioural monitoring or targeted advertising "
        "directed at children.",
        "CHAPTER III Significant Data Fiduciary",
        "13. A Significant Data Fiduciary shall appoint a Data Protection Officer based in India.",
        "14. A Significant Data Fiduciary shall undertake a periodic Data Protection Impact Assessment and an "
        "independent data audit.",
        "CHAPTER IV Rights of Data Principal",
        "15. A Data Fiduciary shall, on request, provide the Data Principal with information about the "
        "personal data processed and the processing activities.",
        "16. A Data Fiduciary shall give effect to a request of the Data Principal for correction, completion, "
        "updating and erasure of personal data.",
        "CHAPTER V Data Protection Board",
        "17. The Central Government shall establish the Data Protection Board of India.",
    ]),
    "RBI-ITO-2023": ("Outsourcing of IT Services Directions - Demo Excerpt", [
        "CHAPTER I Governance",
        "1. Every RE shall put in place a Board approved policy on IT outsourcing.",
        "2. The IT Strategy Committee of the RE shall exercise oversight of all IT outsourcing arrangements.",
        "CHAPTER II Engagement of Service Providers",
        "3. The RE shall assess the materiality of each outsourcing arrangement before entering into it.",
        "4. The RE shall maintain an inventory of all outsourced IT services.",
        "5. The RE shall carry out due diligence of the service provider before outsourcing.",
        "6. The RE shall ensure that the agreement provides access and audit rights to the regulated entity "
        "and to the RBI.",
        "7. The RE shall ensure that the agreement covers confidentiality, sub-contracting and termination.",
        "CHAPTER III Monitoring",
        "8. The RE shall undertake periodic risk assessment and performance monitoring of service providers.",
        "9. The RE shall put in place controls over service provider access to regulated entity systems and data.",
        "10. The RE shall monitor the concentration risk of service providers.",
        "CHAPTER IV Business Continuity and Exit",
        "11. The RE shall ensure that a business continuity and disaster recovery plan with the service "
        "provider is in place and periodically tested.",
        "12. The RE shall have an exit strategy for outsourcing arrangements.",
    ]),
}


def write_all(version: str = "v-demo", overrides: dict[str, str] | None = None,
              only: set[str] | None = None) -> list[Path]:
    paths = []
    for reg_id, (title, lines) in REGS.items():
        if only and reg_id not in only:
            continue
        doc = Document()
        doc.add_heading(title, level=0)
        doc.add_paragraph(DISCLAIMER)
        for line in lines:
            for old, new in (overrides or {}).items():
                line = line.replace(old, new)
            doc.add_paragraph(line)
        out = get_settings().data_dir / "regulations" / reg_id / version / "raw.docx"
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
        paths.append(out)
    return paths


if __name__ == "__main__":
    for p in write_all():
        print("wrote", p)